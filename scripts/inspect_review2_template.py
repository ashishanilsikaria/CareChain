from pathlib import Path
import pptx

ROOT = Path(__file__).resolve().parents[1]
REVIEW2_DIR = ROOT / "docs" / "reviews" / "Review 2"
ppt_path = str(REVIEW2_DIR / "4 Review 2 PPT Template.pptx")

prs = pptx.Presentation(ppt_path)
print(f'Total slides: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides):
    print(f'=== SLIDE {idx+1} ===')
    for s_idx, shape in enumerate(slide.shapes):
        txt = ''
        if shape.has_text_frame:
            txt = ' | '.join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
        print(f'  Shape {s_idx} ({shape.shape_type}, name="{shape.name}"): {txt}')
