from pathlib import Path
import pptx

ROOT = Path(__file__).resolve().parents[1]
REVIEW2_DIR = ROOT / "docs" / "reviews" / "Review 2"
ppt_path = str(REVIEW2_DIR / "CareChain_Review2_Presentation.pptx")

prs = pptx.Presentation(ppt_path)
print(f"Verified Presentation: {len(prs.slides)} slides.")
for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx+1} ---")
    for s in slide.shapes:
        txt = ""
        if s.has_text_frame:
            txt = " // ".join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE] (name={s.name}, size={s.width.inches:.2f}\" x {s.height.inches:.2f}\")")
        elif txt:
            print(f"  [TEXT] {txt[:120]}")
