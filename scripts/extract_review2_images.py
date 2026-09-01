import os
from pathlib import Path
import pptx
import docx

ROOT = Path(__file__).resolve().parents[1]
REVIEW2_DIR = ROOT / "docs" / "reviews" / "Review 2"

ppt_path = str(REVIEW2_DIR / "4 Review 2 PPT Template.pptx")
ppt_img_dir = str(REVIEW2_DIR / "extracted_ppt_images")
os.makedirs(ppt_img_dir, exist_ok=True)

prs = pptx.Presentation(ppt_path)
for idx, slide in enumerate(prs.slides):
    img_count = 0
    for s in slide.shapes:
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            img_count += 1
            image = s.image
            filename = f"slide_{idx+1}_img_{img_count}.{image.ext}"
            filepath = os.path.join(ppt_img_dir, filename)
            with open(filepath, 'wb') as f:
                f.write(image.blob)
            print(f'Saved {filepath}, dim: {s.width.inches:.2f}" x {s.height.inches:.2f}"')

doc_path = str(REVIEW2_DIR / "3 BCSE497J Project I Report - Template.docx")
doc_img_dir = str(REVIEW2_DIR / "extracted_doc_images")
os.makedirs(doc_img_dir, exist_ok=True)

doc = docx.Document(doc_path)
img_idx = 0
for rel in doc.part.rels.values():
    if "image" in rel.target_ref:
        img_idx += 1
        img_part = rel.target_part
        ext = rel.target_ref.split('.')[-1]
        filepath = os.path.join(doc_img_dir, f"doc_img_{img_idx}.{ext}")
        with open(filepath, 'wb') as f:
            f.write(img_part.blob)
        print(f'Saved {filepath}, size: {len(img_part.blob)} bytes')
