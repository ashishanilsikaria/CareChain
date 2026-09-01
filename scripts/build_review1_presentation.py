import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REVIEW1_DIR = ROOT / "docs" / "reviews" / "Review 1"
TEMPLATE = REVIEW1_DIR / "Review 1 - 29 07 2026 format.pptx"
OUTPUT = REVIEW1_DIR / "CareChain_Review1_Presentation.pptx"

REVIEW_DATE = "29-07-2026"
COURSE = "BCSE497J - Project I"
PROGRAMME = "B.Tech Computer Science and Engineering"
SCHOOL = "SCOPE"
STUDENT = "Ashish Anil Sikaria"
REG_NO = "23BCI0218"
GUIDE = "Dr. Selvi M"
PROJECT_NAME = "CareChain"
PROJECT_SUBTITLE = "Decentralized Electronic Health Record Platform"

TITLE_COLOR = RGBColor(31, 78, 121)
BODY_COLOR = RGBColor(40, 40, 40)
MUTED_COLOR = RGBColor(90, 90, 90)


def clear_text(shape):
    if not shape.has_text_frame:
        return
    shape.text_frame.clear()


def set_text(
    shape,
    paragraphs,
    font_size=22,
    bold_first=False,
    align=PP_ALIGN.LEFT,
    space_after=5,
    vertical_anchor=MSO_ANCHOR.TOP,
):
    """Replace a text shape while preserving the template's position."""
    clear_text(shape)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = vertical_anchor

    for index, item in enumerate(paragraphs):
        text, size, bold = item if isinstance(item, tuple) else (item, font_size, False)
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.text = text
        para.alignment = align
        para.space_after = Pt(space_after)
        para.font.name = "Calibri"
        para.font.size = Pt(size)
        para.font.bold = bool(bold or (bold_first and index == 0))
        para.font.color.rgb = TITLE_COLOR if para.font.bold else BODY_COLOR


def set_footer(slide, slide_no):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if text == "21-07-2026":
            set_text(shape, [(REVIEW_DATE, 12, False)], font_size=12, align=PP_ALIGN.LEFT)
        elif text == "SCOPE":
            set_text(shape, [(SCHOOL, 12, False)], font_size=12, align=PP_ALIGN.CENTER)
        elif text == str(slide_no):
            set_text(shape, [(str(slide_no), 12, False)], font_size=12, align=PP_ALIGN.RIGHT)


def style_title(shape, title):
    set_text(shape, [(title, 28, True)], font_size=28, vertical_anchor=MSO_ANCHOR.MIDDLE)


def create_literature_table_image():
    """Render a compact table for the template's existing picture placeholder."""
    width, height = 1600, 324
    header_height = 56
    row_height = (height - header_height) // 4
    column_edges = (0, 260, 540, 1120, width)
    headers = ("Author / Year", "Research Work", "Methodology / Key Finding", "Research Gap")
    rows = (
        ("Azaria et al. (2016)", "MedRec", "Permission management and audit trail", "No encrypted IPFS storage workflow"),
        ("Yue et al. (2016)", "Healthcare Data Gateways", "Purpose-centric patient data access", "No decentralized encrypted storage implementation"),
        ("Xia et al. (2017)", "MeDShare", "Provenance across cloud providers", "No direct patient-to-doctor consent portal"),
        ("Dubovitskaya et al. (2017)", "Secure EMR Sharing", "Fine-grained EMR access control", "No lightweight patient-operated IPFS workflow"),
    )

    image = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    header_font = ImageFont.truetype(str(font_path), 19)
    body_font = ImageFont.truetype(str(font_path), 16)

    draw.rectangle((0, 0, width, header_height), fill=(31, 78, 121))
    for row_index in range(4):
        top = header_height + row_index * row_height
        if row_index % 2 == 0:
            draw.rectangle((0, top, width, top + row_height), fill=(232, 242, 250))

    for edge in column_edges:
        draw.line((edge, 0, edge, height), fill=(110, 130, 150), width=2)
    for y in (0, header_height, height):
        draw.line((0, y, width, y), fill=(110, 130, 150), width=2)
    for row_index in range(1, 4):
        y = header_height + row_index * row_height
        draw.line((0, y, width, y), fill=(150, 165, 180), width=1)

    for column_index, text in enumerate(headers):
        draw.text((column_edges[column_index] + 12, 18), text, font=header_font, fill=(255, 255, 255))
    for row_index, row in enumerate(rows):
        y = header_height + row_index * row_height + 14
        for column_index, text in enumerate(row):
            draw.text((column_edges[column_index] + 12, y), text, font=body_font, fill=(40, 40, 40))

    with tempfile.NamedTemporaryFile(prefix="carechain-literature-", suffix=".png", delete=False) as handle:
        output = Path(handle.name)
    image.save(output)
    return output


def replace_placeholder_with_picture(slide, placeholder_index, image_path):
    """Fill an image placeholder while retaining its template geometry."""
    placeholder = slide.shapes[placeholder_index]
    replacement = slide.shapes.add_picture(
        str(image_path), placeholder.left, placeholder.top, placeholder.width, placeholder.height
    )
    placeholder._element.getparent().remove(placeholder._element)
    return replacement


def main():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    prs = Presentation(str(TEMPLATE))

    # Slide 1: title and academic details.
    slide = prs.slides[0]
    set_text(
        slide.shapes[1],
        [
            (f"{PROGRAMME} | {COURSE}", 15, False),
            (PROJECT_NAME, 27, True),
            (PROJECT_SUBTITLE, 16, False),
        ],
        align=PP_ALIGN.CENTER,
        space_after=0,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    set_text(
        slide.shapes[2],
        [
            ("Student", 16, True),
            (f"{STUDENT} ({REG_NO})", 16, False),
            ("", 8, False),
            ("Faculty Guide", 16, True),
            (GUIDE, 16, False),
            ("Review 1 | Panel 37 | SJT408", 13, False),
        ],
        font_size=16,
        align=PP_ALIGN.CENTER,
        space_after=2,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    # Slide 2: aim.
    slide = prs.slides[1]
    style_title(slide.shapes[1], "Aim")
    set_text(
        slide.shapes[0],
        [
            (
                "To design and implement a secure, patient-centric EHR platform where medical records are encrypted, stored off-chain on IPFS, and governed through blockchain-based consent.",
                20,
                False,
            ),
            (
                "The prototype enables controlled record sharing, verified practitioner access, and an auditable permission history without placing clinical files on-chain.",
                18,
                False,
            ),
        ],
        font_size=20,
        space_after=12,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    # Slide 3: abstract.
    slide = prs.slides[2]
    style_title(slide.shapes[1], "Abstract")
    set_text(
        slide.shapes[0],
        [
            (
                "CareChain addresses privacy, access-control, and auditability limitations in centralized electronic health record systems.",
                18,
                False,
            ),
            (
                "It combines a React role-based portal, Node.js/Express APIs, MongoDB metadata storage, AES-256-GCM encryption, IPFS storage through Pinata, and a Solidity smart contract for access control and audit events.",
                18,
                False,
            ),
            (
                "Patients control sharing; doctors access only authorized records; and administrators verify practitioners. Encrypted record references and access events are recorded on-chain, while clinical files remain off-chain.",
                18,
                False,
            ),
        ],
        font_size=18,
        space_after=9,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    # Slide 4: literature review.
    slide = prs.slides[3]
    style_title(slide.shapes[1], "Literature Review")
    table_image = create_literature_table_image()
    replace_placeholder_with_picture(slide, 0, table_image)
    table_image.unlink(missing_ok=True)

    # Slide 5: research gap.
    slide = prs.slides[4]
    style_title(slide.shapes[1], "Research Gap")
    set_text(
        slide.shapes[0],
        [
            ("Centralized EHR repositories create a single high-value target for sensitive medical data.", 17, False),
            ("Blockchain-only approaches are unsuitable for storing clinical files because of privacy and cost constraints.", 17, False),
            ("IPFS storage requires encryption, key management, and a clear patient-controlled sharing workflow.", 17, False),
            ("CareChain combines encrypted off-chain records with on-chain permissions and separate patient, doctor, and administrator portals.", 17, True),
        ],
        font_size=17,
        space_after=10,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )

    # Slide 6: objectives, SDG, TRL and expected outcome.
    slide = prs.slides[5]
    style_title(slide.shapes[1], "Objectives")
    set_text(
        slide.shapes[0],
        [
            ("1. Build a patient portal for encrypted record upload, viewing, and selective sharing.", 15, False),
            ("2. Implement doctor and administrator workflows for authorized access and practitioner verification.", 15, False),
            ("3. Use a Solidity smart contract to manage permissions and produce immutable access events.", 15, False),
            ("4. Keep clinical files off-chain through AES-256-GCM encryption and IPFS storage.", 15, False),
        ],
        font_size=15,
        space_after=5,
    )
    set_text(
        slide.shapes[5],
        [
            ("Identified SDG: SDG 3 - Good Health and Well-being", 11, True),
            ("TRL 4: Technology validated in a laboratory environment", 11, True),
            ("Outcome: Functional EHR prototype, source code, and test report", 11, False),
        ],
        font_size=11,
        align=PP_ALIGN.LEFT,
        space_after=0,
    )

    # Slide 7: references.
    slide = prs.slides[6]
    style_title(slide.shapes[0], "References")
    set_text(
        slide.shapes[1],
        [
            ('[1] A. Azaria et al., "MedRec: Using blockchain for medical data access and permission management," 2016 IEEE 2nd International Conference on Open and Big Data, pp. 25-30.', 11, False),
            ('[2] X. Yue et al., "Healthcare Data Gateways: Found Healthcare Intelligence on Blockchain with Novel Privacy Risk Control," Journal of Medical Systems, vol. 40, no. 10, art. 218, 2016.', 11, False),
            ('[3] Q. Xia et al., "MeDShare: Trust-Less Medical Data Sharing among Cloud Service Providers via Blockchain," IEEE Access, vol. 5, pp. 14757-14767, 2017.', 11, False),
            ('[4] A. Dubovitskaya et al., "Secure and Trustable Electronic Medical Records Sharing using Blockchain," AMIA Annual Symposium Proceedings, pp. 650-659, 2017.', 11, False),
        ],
        font_size=11,
        space_after=4,
    )

    for index, slide in enumerate(prs.slides, start=1):
        set_footer(slide, index)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Generated {OUTPUT}")


if __name__ == "__main__":
    main()
